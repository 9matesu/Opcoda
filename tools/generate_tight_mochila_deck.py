import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_tight_mochila_deck(output_path):
    prs = Presentation()
    # Exact Canva widescreen standard (20.0 x 11.25 inches)
    prs.slide_width = Inches(20.0)
    prs.slide_height = Inches(11.25)

    # Color tokens from OpcodaProg Linear.pptx
    BG_COLOR = RGBColor(0xF4, 0xF5, 0xF6)       # Clean off-white
    SURFACE_COLOR = RGBColor(0xFF, 0xFF, 0xFF)  # Pure white
    TEXT_COLOR = RGBColor(0x12, 0x15, 0x18)     # Near Black
    MUTED_COLOR = RGBColor(0x5A, 0x60, 0x65)    # Slate / Gray
    ACCENT_COLOR = RGBColor(0xFF, 0x9A, 0x00)   # Opcoda Amber (#FF9A00)
    BORDER_COLOR = RGBColor(0xD8, 0xDC, 0xDE)   # Light divider
    DARK_BOX = RGBColor(0x0E, 0x10, 0x13)       # Terminal Dark
    HIGHLIGHT_BG = RGBColor(0xFF, 0xFD, 0xF9)   # Warm highlight
    GREEN_ACCENT = RGBColor(0x2E, 0x7D, 0x32)   # Green / Accepted
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    blank_layout = prs.slide_layouts[6]
    MAIN_SLIDES = 13

    def set_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, number_str, section_str, speaker_str=""):
        # Slide number
        num_box = slide.shapes.add_textbox(Inches(1.3), Inches(0.7), Inches(1.4), Inches(0.6))
        tf = num_box.text_frame
        p = tf.paragraphs[0]
        p.text = number_str
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR

        # Section tag
        sec_box = slide.shapes.add_textbox(Inches(2.5), Inches(0.72), Inches(11.0), Inches(0.6))
        tf = sec_box.text_frame
        p = tf.paragraphs[0]
        p.text = section_str.upper()
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = MUTED_COLOR

        # Top divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(1.45), Inches(17.41), Inches(0.025))
        line.fill.solid()
        line.fill.fore_color.rgb = TEXT_COLOR
        line.line.fill.background()

    def add_footer(slide, number_str, is_appendix=False):
        # Footer divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(10.15), Inches(17.41), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.fill.background()

        # Left tag
        fb = slide.shapes.add_textbox(Inches(1.3), Inches(10.25), Inches(10.0), Inches(0.5))
        p = fb.text_frame.paragraphs[0]
        p.text = "OPCODA · TÊMPERA SIMULADA NA MOCHILA" if not is_appendix else "OPCODA · MATERIAL COMPLEMENTAR (BACKUP)"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = MUTED_COLOR

        # Right count
        fb_r = slide.shapes.add_textbox(Inches(15.7), Inches(10.25), Inches(3.0), Inches(0.5))
        p = fb_r.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"{number_str} / {MAIN_SLIDES:02d}" if not is_appendix else f"BACKUP {number_str}"
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = MUTED_COLOR

    def add_title(slide, title_text):
        tb = slide.shapes.add_textbox(Inches(1.3), Inches(1.85), Inches(17.41), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = TEXT_COLOR

    # =============================================================
    # SLIDE 01 — CAPA
    # =============================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_background(s1)
    add_header(s1, "01", "FATEC Cruzeiro · Programação Linear / Otimização")

    tag1 = s1.shapes.add_textbox(Inches(1.3), Inches(2.4), Inches(15.0), Inches(0.6))
    p = tag1.text_frame.paragraphs[0]
    p.text = "TRABALHO ACADÊMICO DE OTIMIZAÇÃO · PROJETO OPCODA"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    t1 = s1.shapes.add_textbox(Inches(1.3), Inches(3.0), Inches(17.4), Inches(2.2))
    tf1 = t1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Aplicação de Têmpera Simulada\nno Problema da Mochila"
    p1.font.name = "Arial"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_COLOR

    sub1 = s1.shapes.add_textbox(Inches(1.3), Inches(5.4), Inches(17.4), Inches(1.0))
    p_sub = sub1.text_frame.paragraphs[0]
    p_sub.text = "Uma aplicação experimental no contexto do Opcoda"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(26)
    p_sub.font.color.rgb = MUTED_COLOR

    # Card autores
    card1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(6.8), Inches(10.5), Inches(2.2))
    card1.fill.solid()
    card1.fill.fore_color.rgb = SURFACE_COLOR
    card1.line.color.rgb = BORDER_COLOR
    card1.line.width = Pt(1)

    bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(6.8), Inches(0.18), Inches(2.2))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = ACCENT_COLOR
    bar1.line.fill.background()

    aut1 = s1.shapes.add_textbox(Inches(1.8), Inches(7.05), Inches(9.5), Inches(1.8))
    p_aut = aut1.text_frame.paragraphs[0]
    p_aut.text = "Mateus Costa · Rafael Elisei"
    p_aut.font.name = "Arial"
    p_aut.font.size = Pt(24)
    p_aut.font.bold = True
    p_aut.font.color.rgb = TEXT_COLOR
    p_aut2 = aut1.text_frame.add_paragraph()
    p_aut2.text = "5º semestre · Análise e Desenvolvimento de Sistemas\nFATEC Cruzeiro"
    p_aut2.font.name = "Arial"
    p_aut2.font.size = Pt(18)
    p_aut2.font.color.rgb = MUTED_COLOR
    p_aut2.space_before = Pt(8)

    add_footer(s1, "01")

    # =============================================================
    # SLIDE 02 — OPCODA (PIPELINE VISUAL LIMPO)
    # =============================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_background(s2)
    add_header(s2, "02", "Fundamentação · Contexto")
    add_title(s2, "O que é o Opcoda?")

    # 3 bullets diretos
    tb_b2 = s2.shapes.add_textbox(Inches(1.3), Inches(3.2), Inches(17.4), Inches(2.2))
    tf2 = tb_b2.text_frame
    bullets2 = [
        "Instrumento virtual experimental em tempo real",
        "Converte dados de arquivos executáveis (.exe) em material sonoro",
        "Utiliza síntese granular para a geração acústica"
    ]
    for i, b in enumerate(bullets2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"■   {b}"
        p.font.name = "Arial"
        p.font.size = Pt(25)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(22)

    # Pipeline visual horizontal em blocos
    nodes2 = [(".EXE", 2.0), ("BYTES", 2.2), ("BLOCOS", 2.3), ("GRÃOS", 2.2), ("SOM", 2.0)]
    x = 1.3
    for i, (nom, w_val) in enumerate(nodes2):
        is_last = (i == len(nodes2) - 1)
        box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(6.8), Inches(w_val), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = ACCENT_COLOR if is_last else SURFACE_COLOR
        box.line.color.rgb = ACCENT_COLOR if is_last else TEXT_COLOR
        box.line.width = Pt(2)
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = nom
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = WHITE if is_last else TEXT_COLOR
        x += w_val

        if not is_last:
            seta = s2.shapes.add_textbox(Inches(x), Inches(7.1), Inches(0.9), Inches(0.9))
            p = seta.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = "→"
            p.font.size = Pt(30)
            p.font.bold = True
            p.font.color.rgb = ACCENT_COLOR
            x += 0.9

    add_footer(s2, "02")

    # =============================================================
    # SLIDE 03 — O PROBLEMA (REDUÇÃO AGRESSIVA E DIRETA)
    # =============================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_background(s3)
    add_header(s3, "03", "Motivação · O Problema")
    add_title(s3, "Onde entra a otimização?")

    # 4 frases de impacto direto em cards amplos
    prob_lines = [
        "O arquivo gera muitos blocos de dados.",
        "Nem todos são interessantes para a síntese sonora.",
        "Existe uma capacidade limitada de processamento e tempo.",
        "Precisamos escolher os melhores blocos."
    ]

    for i, line in enumerate(prob_lines):
        is_key = (i == 3)
        c_box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2 + i * 1.55), Inches(17.41), Inches(1.25))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = HIGHLIGHT_BG if is_key else SURFACE_COLOR
        c_box.line.color.rgb = ACCENT_COLOR if is_key else BORDER_COLOR
        c_box.line.width = Pt(2 if is_key else 1)

        b_bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2 + i * 1.55), Inches(0.16), Inches(1.25))
        b_bar.fill.solid()
        b_bar.fill.fore_color.rgb = ACCENT_COLOR
        b_bar.line.fill.background()

        tb = s3.shapes.add_textbox(Inches(1.8), Inches(3.45 + i * 1.55), Inches(16.5), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(26)
        p.font.bold = is_key
        p.font.color.rgb = ACCENT_COLOR if is_key else TEXT_COLOR

    add_footer(s3, "03")

    # =============================================================
    # SLIDE 04 — DO BYTE AO ITEM (CONEXÃO INTUITIVA COM OS BYTES)
    # =============================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_background(s4)
    add_header(s4, "04", "Conceito Central · Do Byte ao Item")
    add_title(s4, "Do byte ao item")

    # Topo: Binário .EXE
    top_c = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.0), Inches(17.41), Inches(1.3))
    top_c.fill.solid()
    top_c.fill.fore_color.rgb = DARK_BOX
    p = top_c.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "ARQUIVO .EXE (DADOS BINÁRIOS)   →   [ Bloco A ]  [ Bloco B ]  [ Bloco C ]  [ Bloco D ]  [ Bloco E ]"
    p.font.name = "Consolas"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # Meio: barras visuais de relevância
    grid_blocos = [
        ("BLOCO A", "█████████████", "interessante (alto valor)", True),
        ("BLOCO B", "███", "pouco relevante", False),
        ("BLOCO C", "████████████████", "interessante (alto valor)", True),
        ("BLOCO D", "██", "pouco relevante", False),
        ("BLOCO E", "███████████", "interessante (alto valor)", True)
    ]

    for i, (bl_nome, barra, status, sel) in enumerate(grid_blocos):
        bx = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(4.6 + i * 0.82), Inches(17.41), Inches(0.72))
        bx.fill.solid()
        bx.fill.fore_color.rgb = HIGHLIGHT_BG if sel else SURFACE_COLOR
        bx.line.color.rgb = ACCENT_COLOR if sel else BORDER_COLOR
        bx.line.width = Pt(1.5 if sel else 1)

        tb = bx.text_frame
        p = tb.paragraphs[0]
        p.text = f"{bl_nome}    {barra}    →   {status}"
        p.font.name = "Consolas"
        p.font.size = Pt(18)
        p.font.bold = sel
        p.font.color.rgb = ACCENT_COLOR if sel else MUTED_COLOR

    # Rodapé do slide: Conexão com Têmpera Simulada
    res_box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(8.9), Inches(17.41), Inches(1.0))
    res_box.fill.solid()
    res_box.fill.fore_color.rgb = SURFACE_COLOR
    res_box.line.color.rgb = TEXT_COLOR
    res_box.line.width = Pt(2)
    p = res_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "TÊMPERA SIMULADA SELECIONA:    BLOCO A  +  BLOCO C  +  BLOCO E"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    add_footer(s4, "04")

    # =============================================================
    # SLIDE 05 — PROBLEMA DA MOCHILA (VISUAL PURO)
    # =============================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_background(s5)
    add_header(s5, "05", "Modelagem · Analogia")
    add_title(s5, "Problema da Mochila")

    # Blocos disponíveis
    lbl_b = s5.shapes.add_textbox(Inches(1.3), Inches(3.0), Inches(17.41), Inches(0.5))
    p = lbl_b.text_frame.paragraphs[0]
    p.text = "BLOCOS CANDIDATOS DISPONÍVEIS:"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR

    blocos_letras = ["A", "B", "C", "D", "E"]
    for i, let in enumerate(blocos_letras):
        is_sel = let in ["A", "C", "E"]
        bx = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3 + i * 3.6), Inches(3.6), Inches(3.0), Inches(1.3))
        bx.fill.solid()
        bx.fill.fore_color.rgb = HIGHLIGHT_BG if is_sel else SURFACE_COLOR
        bx.line.color.rgb = ACCENT_COLOR if is_sel else BORDER_COLOR
        bx.line.width = Pt(2 if is_sel else 1)
        p = bx.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f"BLOCO {let}"
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR if is_sel else TEXT_COLOR

    # Seta
    st = s5.shapes.add_textbox(Inches(9.5), Inches(5.1), Inches(1.0), Inches(0.7))
    p = st.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "↓"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    # A Mochila
    mochila_card = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(5.8), Inches(11.0), Inches(2.3))
    mochila_card.fill.solid()
    mochila_card.fill.fore_color.rgb = SURFACE_COLOR
    mochila_card.line.color.rgb = ACCENT_COLOR
    mochila_card.line.width = Pt(2.5)

    tb_m = mochila_card.text_frame
    p = tb_m.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "MOCHILA (CAPACIDADE LIMITADA = C)"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR

    p2 = tb_m.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "[  BLOCO A    ·    BLOCO C    ·    BLOCO E  ]"
    p2.font.name = "Arial"
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_COLOR
    p2.space_before = Pt(14)

    # Frase final
    fr = s5.shapes.add_textbox(Inches(1.3), Inches(8.7), Inches(17.41), Inches(1.0))
    p = fr.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "“Selecionar blocos de maior relevância dentro de uma capacidade limitada.”"
    p.font.name = "Arial"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR

    add_footer(s5, "05")

    # =============================================================
    # SLIDE 06 — COMO UMA SOLUÇÃO É REPRESENTADA
    # =============================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_background(s6)
    add_header(s6, "06", "Modelagem · Vetor Binário")
    add_title(s6, "Como a solução é representada?")

    # Vetor de blocos e vetor binário
    v_box = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2), Inches(17.41), Inches(2.6))
    v_box.fill.solid()
    v_box.fill.fore_color.rgb = SURFACE_COLOR
    v_box.line.color.rgb = TEXT_COLOR
    v_box.line.width = Pt(2)

    tb6 = v_box.text_frame
    p = tb6.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "BLOCOS:       [    A   ,    B   ,    C   ,    D   ,    E    ]"
    p.font.name = "Consolas"
    p.font.size = Pt(28)
    p.font.color.rgb = MUTED_COLOR

    p2 = tb6.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "SOLUÇÃO:      [    1   ,    0   ,    1   ,    0   ,    1    ]"
    p2.font.name = "Consolas"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_COLOR
    p2.space_before = Pt(14)

    # Explicação limpa
    tb_exp6 = s6.shapes.add_textbox(Inches(1.3), Inches(6.3), Inches(17.41), Inches(2.2))
    tf6 = tb_exp6.text_frame
    bullets6 = [
        "1  →  bloco selecionado para a síntese",
        "0  →  bloco descartado",
        "Restrição: Peso Total ≤ Capacidade (C)"
    ]
    for i, b in enumerate(bullets6):
        p = tf6.paragraphs[0] if i == 0 else tf6.add_paragraph()
        p.text = f"■   {b}"
        p.font.name = "Arial"
        p.font.size = Pt(25)
        p.font.bold = (i == 2)
        p.font.color.rgb = ACCENT_COLOR if (i == 2) else TEXT_COLOR
        p.space_after = Pt(16)

    add_footer(s6, "06")

    # =============================================================
    # SLIDE 07 — TÊMPERA SIMULADA
    # =============================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_background(s7)
    add_header(s7, "07", "Metaheurística · Conceito")
    add_title(s7, "O que é Têmpera Simulada?")

    # 3 bullets
    tb7 = s7.shapes.add_textbox(Inches(1.3), Inches(3.2), Inches(17.41), Inches(2.8))
    tf7 = tb7.text_frame
    bullets7 = [
        "Metaheurística de otimização inspirada no resfriamento de materiais",
        "Explora diferentes combinações de blocos no espaço de soluções",
        "Pode aceitar soluções piores temporariamente ao longo da busca"
    ]
    for i, b in enumerate(bullets7):
        p = tf7.paragraphs[0] if i == 0 else tf7.add_paragraph()
        p.text = f"■   {b}"
        p.font.name = "Arial"
        p.font.size = Pt(25)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(20)

    # Grande card de destaque
    card7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(6.5), Inches(17.41), Inches(2.4))
    card7.fill.solid()
    card7.fill.fore_color.rgb = HIGHLIGHT_BG
    card7.line.color.rgb = ACCENT_COLOR
    card7.line.width = Pt(2)

    bar7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(6.5), Inches(0.2), Inches(2.4))
    bar7.fill.solid()
    bar7.fill.fore_color.rgb = ACCENT_COLOR
    bar7.line.fill.background()

    tb_dest7 = s7.shapes.add_textbox(Inches(1.8), Inches(6.9), Inches(16.5), Inches(1.8))
    p = tb_dest7.text_frame.paragraphs[0]
    p.text = "“A Têmpera Simulada permite aceitar temporariamente uma solução pior para escapar de ótimos locais.”"
    p.font.name = "Arial"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR

    p2 = tb_dest7.text_frame.add_paragraph()
    p2.text = "Kirkpatrick, Gelatt e Vecchi (1983)"
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = MUTED_COLOR
    p2.space_before = Pt(12)

    add_footer(s7, "07")

    # =============================================================
    # SLIDE 08 — COMO O ALGORITMO PROCURA (SLIDE ÚNICO E VISUAL)
    # =============================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_background(s8)
    add_header(s8, "08", "Funcionamento · Vizinhança")
    add_title(s8, "Como o algoritmo procura?")

    # Fluxo visual em cards verticais
    c_s = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2), Inches(17.41), Inches(1.4))
    c_s.fill.solid()
    c_s.fill.fore_color.rgb = SURFACE_COLOR
    c_s.line.color.rgb = BORDER_COLOR
    c_s.line.width = Pt(1.5)
    p = c_s.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "SOLUÇÃO ATUAL:      [  1 ,   0 ,   1 ,   0 ,   1  ]"
    p.font.name = "Consolas"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR

    st8 = s8.shapes.add_textbox(Inches(9.5), Inches(4.7), Inches(1.0), Inches(0.7))
    p = st8.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "↓  inverte um bloco (bit flip)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    c_s_nova = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(5.5), Inches(17.41), Inches(1.4))
    c_s_nova.fill.solid()
    c_s_nova.fill.fore_color.rgb = HIGHLIGHT_BG
    c_s_nova.line.color.rgb = ACCENT_COLOR
    c_s_nova.line.width = Pt(2)
    p = c_s_nova.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "NOVA SOLUÇÃO:       [  1 ,   1 ,   1 ,   0 ,   1  ]"
    p.font.name = "Consolas"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    # Decisão abaixo
    c_dec = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(7.3), Inches(17.41), Inches(2.0))
    c_dec.fill.solid()
    c_dec.fill.fore_color.rgb = SURFACE_COLOR
    c_dec.line.color.rgb = TEXT_COLOR
    c_dec.line.width = Pt(2)

    tb_dec = c_dec.text_frame
    p = tb_dec.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "A NOVA SOLUÇÃO É MELHOR?"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR

    p2 = tb_dec.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "SIM  →  ACEITA IMEDIATAMENTE          |          NÃO  →  PODE ACEITAR (PROBABILIDADE)"
    p2.font.name = "Arial"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_COLOR
    p2.space_before = Pt(12)

    add_footer(s8, "08")

    # =============================================================
    # SLIDE 09 — POR QUE ACEITAR UMA SOLUÇÃO PIOR? (A FÓRMULA ÚNICA)
    # =============================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_background(s9)
    add_header(s9, "09", "Critério de Metropolis")
    add_title(s9, "Por que aceitar uma solução pior?")

    # A única fórmula em tamanho grande
    f9 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2), Inches(8.0), Inches(2.2))
    f9.fill.solid()
    f9.fill.fore_color.rgb = SURFACE_COLOR
    f9.line.color.rgb = ACCENT_COLOR
    f9.line.width = Pt(2.5)

    tb_f9 = f9.text_frame
    p = tb_f9.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "P  =  e^( Δ / T )"
    p.font.name = "Arial"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    p2 = tb_f9.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Probabilidade de aceitar piora (Metropolis)"
    p2.font.size = Pt(16)
    p2.font.color.rgb = MUTED_COLOR
    p2.space_before = Pt(8)

    # Exemplo visual ao lado
    ex9 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.8), Inches(3.2), Inches(8.91), Inches(2.2))
    ex9.fill.solid()
    ex9.fill.fore_color.rgb = SURFACE_COLOR
    ex9.line.color.rgb = BORDER_COLOR
    ex9.line.width = Pt(1.5)

    tb_ex = ex9.text_frame
    p = tb_ex.paragraphs[0]
    p.text = "EXEMPLO PRÁTICO:"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR

    p2 = tb_ex.add_paragraph()
    p2.text = "Solução atual = 100    →    Nova solução = 90  (Piorou)"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_COLOR
    p2.space_before = Pt(6)

    # Dois casos abaixo
    c_alta = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(5.9), Inches(17.41), Inches(1.6))
    c_alta.fill.solid()
    c_alta.fill.fore_color.rgb = HIGHLIGHT_BG
    c_alta.line.color.rgb = ACCENT_COLOR
    c_alta.line.width = Pt(2)
    p = c_alta.text_frame.paragraphs[0]
    p.text = "TEMPERATURA ALTA (início):"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p2 = c_alta.text_frame.add_paragraph()
    p2.text = "Ainda pode aceitar a solução pior para explorar novas combinações."
    p2.font.size = Pt(22)
    p2.font.color.rgb = TEXT_COLOR
    p2.space_before = Pt(4)

    c_baixa = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(7.8), Inches(17.41), Inches(1.6))
    c_baixa.fill.solid()
    c_baixa.fill.fore_color.rgb = SURFACE_COLOR
    c_baixa.line.color.rgb = BORDER_COLOR
    c_baixa.line.width = Pt(1.5)
    p = c_baixa.text_frame.paragraphs[0]
    p.text = "TEMPERATURA BAIXA (fim):"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR
    p2 = c_baixa.text_frame.add_paragraph()
    p2.text = "Provavelmente rejeita a solução pior, tornando a busca estritamente seletiva."
    p2.font.size = Pt(22)
    p2.font.color.rgb = TEXT_COLOR
    p2.space_before = Pt(4)

    add_footer(s9, "09")

    # =============================================================
    # SLIDE 10 — TEMPERATURA + FATOR DE REDUÇÃO (SLIDE ÚNICO E VISUAL)
    # =============================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_background(s10)
    add_header(s10, "10", "Parâmetros · Resfriamento Térmico")
    add_title(s10, "Temperatura e resfriamento")

    # Gráfico esquemático à esquerda
    g10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2), Inches(8.0), Inches(4.8))
    g10.fill.solid()
    g10.fill.fore_color.rgb = DARK_BOX
    g10.line.fill.background()

    tb_g = g10.text_frame
    p = tb_g.paragraphs[0]
    p.text = "CURVA DE RESFRIAMENTO\n"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR

    linhas_grafico = [
        "100 ───────┐",
        "           │  alta exploração",
        " 50 ───    │",
        "           │",
        " 10 ──     │  alta seleção",
        "           │",
        "  0 ───────┘"
    ]
    for lg in linhas_grafico:
        p = tb_g.add_paragraph()
        p.text = lg
        p.font.name = "Consolas"
        p.font.size = Pt(20)
        p.font.color.rgb = WHITE

    # Cards de parâmetros à direita
    params10 = [
        ("T₀", "Começo da busca", "Temperatura inicial alta (ampla exploração)"),
        ("Tmin", "Fim da busca", "Temperatura final (critério de parada)"),
        ("α", "Velocidade do resfriamento", "Fator de redução geométrico:  Tnovo = α · Tatual")
    ]

    for i, (sig, rot, desc) in enumerate(params10):
        bx = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.8), Inches(3.2 + i * 1.65), Inches(8.91), Inches(1.4))
        bx.fill.solid()
        bx.fill.fore_color.rgb = SURFACE_COLOR
        bx.line.color.rgb = BORDER_COLOR
        bx.line.width = Pt(1.5)

        tb = bx.text_frame
        p = tb.paragraphs[0]
        p.text = f"{sig}   ·   {rot}"
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR

        p2 = tb.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(17)
        p2.font.color.rgb = TEXT_COLOR
        p2.space_before = Pt(4)

    # Frase resumo embaixo
    res10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(8.4), Inches(17.41), Inches(1.3))
    res10.fill.solid()
    res10.fill.fore_color.rgb = HIGHLIGHT_BG
    res10.line.color.rgb = ACCENT_COLOR
    res10.line.width = Pt(2)
    p = res10.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Alta temperatura → mais liberdade para explorar         |         Baixa temperatura → busca mais seletiva"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR

    add_footer(s10, "10")

    # =============================================================
    # SLIDE 11 — FLUXO COMPLETO (PIPELINE INTEGRADO)
    # =============================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_background(s11)
    add_header(s11, "11", "Arquitetura · Visão Completa")
    add_title(s11, "Fluxo completo")

    # Blocos sequenciais em grid 2 linhas ou vertical claro
    pipe_nodes = [
        ("1. .EXE", "Arquivo executável"),
        ("2. BLOCOS", "Dados particionados"),
        ("3. CARACTERÍSTICAS", "Relevância / peso"),
        ("4. MOCHILA", "Restrição de capacidade"),
        ("5. TÊMPERA SIMULADA", "Otimização metaheurística"),
        ("6. BLOCOS SELECIONADOS", "Melhor combinação"),
        ("7. SÍNTESE GRANULAR", "Geração de áudio final")
    ]

    # Grid 4 + 3
    for i in range(4):
        nom, sub = pipe_nodes[i]
        bx = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3 + i * 4.45), Inches(3.3), Inches(4.1), Inches(2.2))
        bx.fill.solid()
        bx.fill.fore_color.rgb = SURFACE_COLOR
        bx.line.color.rgb = BORDER_COLOR
        bx.line.width = Pt(1.5)
        p = bx.text_frame.paragraphs[0]
        p.text = nom
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = TEXT_COLOR
        p2 = bx.text_frame.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(16)
        p2.font.color.rgb = MUTED_COLOR
        p2.space_before = Pt(8)

    for i in range(3):
        nom, sub = pipe_nodes[i + 4]
        is_opt = (i == 0)
        bx = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.8 + i * 4.8), Inches(6.0), Inches(4.4), Inches(2.2))
        bx.fill.solid()
        bx.fill.fore_color.rgb = HIGHLIGHT_BG if is_opt else SURFACE_COLOR
        bx.line.color.rgb = ACCENT_COLOR if is_opt else BORDER_COLOR
        bx.line.width = Pt(2.5 if is_opt else 1.5)
        p = bx.text_frame.paragraphs[0]
        p.text = nom
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR if is_opt else TEXT_COLOR
        p2 = bx.text_frame.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(16)
        p2.font.color.rgb = TEXT_COLOR if is_opt else MUTED_COLOR
        p2.space_before = Pt(8)

    c11_f = s11.shapes.add_textbox(Inches(1.3), Inches(8.8), Inches(17.41), Inches(0.8))
    p = c11_f.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "A Têmpera Simulada atua como filtro seletivo ótimo antes da síntese granular."
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR

    add_footer(s11, "11")

    # =============================================================
    # SLIDE 12 — COMO AVALIAR
    # =============================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_background(s12)
    add_header(s12, "12", "Metodologia Experimental")
    add_title(s12, "Como avaliar?")

    tb12 = s12.shapes.add_textbox(Inches(1.3), Inches(3.2), Inches(17.41), Inches(3.5))
    tf12 = tb12.text_frame
    bullets12 = [
        "Comparar o valor da solução inicial com a solução otimizada",
        "Verificar o aproveitamento da capacidade disponível da mochila",
        "Acompanhar o número de iterações até a estabilização do valor",
        "Avaliar a sensibilidade a diferentes velocidades de resfriamento (α)"
    ]
    for i, b in enumerate(bullets12):
        p = tf12.paragraphs[0] if i == 0 else tf12.add_paragraph()
        p.text = f"■   {b}"
        p.font.name = "Arial"
        p.font.size = Pt(25)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(22)

    c12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(7.2), Inches(17.41), Inches(1.8))
    c12.fill.solid()
    c12.fill.fore_color.rgb = SURFACE_COLOR
    c12.line.color.rgb = BORDER_COLOR
    c12.line.width = Pt(1.5)
    p = c12.text_frame.paragraphs[0]
    p.text = "Etapa experimental a ser realizada."
    p.font.name = "Arial"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR
    p2 = c12.text_frame.add_paragraph()
    p2.text = "Apresentamos o modelo e a arquitetura; os dados numéricos serão medidos nos testes com o Opcoda."
    p2.font.size = Pt(18)
    p2.font.color.rgb = MUTED_COLOR
    p2.space_before = Pt(8)

    add_footer(s12, "12")

    # =============================================================
    # SLIDE 13 — CONCLUSÃO / ENCERRAMENTO
    # =============================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_background(s13)
    add_header(s13, "13", "Conclusão · Agradecimentos")

    t13 = s13.shapes.add_textbox(Inches(1.3), Inches(2.2), Inches(17.41), Inches(1.2))
    p = t13.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Opcoda + Têmpera Simulada"
    p.font.name = "Arial"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR

    # Equação visual limpa
    eq_items = [
        ("DADOS BINÁRIOS", False),
        ("+", None),
        ("PROBLEMA DA MOCHILA", False),
        ("+", None),
        ("TÊMPERA SIMULADA", False),
        ("+", None),
        ("SÍNTESE GRANULAR", False),
        ("=", None),
        ("OPCODA", True)
    ]
    x = 1.3
    for label, is_acc in eq_items:
        if is_acc is None:
            tb = s13.shapes.add_textbox(Inches(x), Inches(4.2), Inches(0.6), Inches(1.0))
            p = tb.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = label
            p.font.name = "Arial"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = ACCENT_COLOR
            x += 0.65
        else:
            w_box = 2.8 if "MOCHILA" in label or "TÊMPERA" in label else (2.4 if "DADOS" in label or "SÍNTESE" in label else 2.0)
            bx = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(4.1), Inches(w_box), Inches(1.2))
            bx.fill.solid()
            bx.fill.fore_color.rgb = ACCENT_COLOR if is_acc else SURFACE_COLOR
            bx.line.color.rgb = ACCENT_COLOR if is_acc else TEXT_COLOR
            bx.line.width = Pt(1.5)
            p = bx.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = label
            p.font.name = "Arial"
            p.font.size = Pt(15)
            p.font.bold = True
            p.font.color.rgb = WHITE if is_acc else TEXT_COLOR
            x += w_box + 0.15

    # Card autores final
    c13 = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(6.5), Inches(11.0), Inches(2.2))
    c13.fill.solid()
    c13.fill.fore_color.rgb = SURFACE_COLOR
    c13.line.color.rgb = TEXT_COLOR
    c13.line.width = Pt(2)
    tb13 = c13.text_frame
    p = tb13.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Mateus Costa · Rafael Elisei"
    p.font.name = "Arial"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR
    p2 = tb13.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "5º semestre · ADS · FATEC Cruzeiro · Programação Linear / Otimização"
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.color.rgb = MUTED_COLOR
    p2.space_before = Pt(8)

    add_footer(s13, "13")

    # =============================================================
    # SLIDES RESERVAS / BACKUP (PARA CASO O PROFESSOR PERGUNTE)
    # =============================================================
    
    # RESERVA 1: FORMULAÇÃO MATEMÁTICA FORMAL
    s_b1 = prs.slides.add_slide(blank_layout)
    set_background(s_b1)
    add_header(s_b1, "A1", "Slide Reserva · Formulação Matemática Completa")
    add_title(s_b1, "Formulação matemática formal")
    f_box = s_b1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2), Inches(17.41), Inches(2.8))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = SURFACE_COLOR
    f_box.line.color.rgb = TEXT_COLOR
    f_box.line.width = Pt(2)
    tf_f = f_box.text_frame
    p = tf_f.paragraphs[0]
    p.text = "max   ∑ (vᵢ · xᵢ)"
    p.font.name = "Arial"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p2 = tf_f.add_paragraph()
    p2.text = "sujeito a:   ∑ (wᵢ · xᵢ)  ≤  C"
    p2.font.name = "Arial"
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_COLOR
    p2.space_before = Pt(12)
    p3 = tf_f.add_paragraph()
    p3.text = "xᵢ  ∈  { 0, 1 }      para todo i = 1, ..., n"
    p3.font.name = "Arial"
    p3.font.size = Pt(24)
    p3.font.color.rgb = MUTED_COLOR
    p3.space_before = Pt(12)
    add_footer(s_b1, "01", is_appendix=True)

    # RESERVA 2: PSEUDOCÓDIGO COMPLETO
    s_b2 = prs.slides.add_slide(blank_layout)
    set_background(s_b2)
    add_header(s_b2, "A2", "Slide Reserva · Pseudocódigo Completo")
    add_title(s_b2, "Pseudocódigo completo do algoritmo")
    code_bg = s_b2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2), Inches(17.41), Inches(6.5))
    code_bg.fill.solid()
    code_bg.fill.fore_color.rgb = DARK_BOX
    code_bg.line.fill.background()
    tb_c = s_b2.shapes.add_textbox(Inches(1.8), Inches(3.4), Inches(16.5), Inches(6.1))
    tf_c = tb_c.text_frame
    pseudo_lines = [
        "GERAR solução inicial viável S (peso ≤ C)",
        "AVALIAR solução inicial: V(S)",
        "T = T₀",
        "ENQUANTO T > Tmin:",
        "    GERAR vizinha S' por mutação (bit flip)",
        "    SE S' viável (peso ≤ C):",
        "        Δ = V(S') − V(S)",
        "        SE Δ > 0: S = S'   // aceita melhoria",
        "        SENÃO SE random(0, 1) < exp(Δ / T): S = S'   // aceita piora",
        "    REDUZIR temperatura: T = α · T",
        "RETORNAR melhor solução encontrada"
    ]
    for i, line in enumerate(pseudo_lines):
        p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(20)
        p.font.color.rgb = ACCENT_COLOR if any(k in line for k in ["GERAR", "ENQUANTO", "SE", "RETORNAR"]) else WHITE
        p.space_after = Pt(6)
    add_footer(s_b2, "02", is_appendix=True)

    # RESERVA 3: REFERÊNCIAS BIBLIOGRÁFICAS
    s_b3 = prs.slides.add_slide(blank_layout)
    set_background(s_b3)
    add_header(s_b3, "A3", "Slide Reserva · Referências Acadêmicas")
    add_title(s_b3, "Referências bibliográficas")
    refs = [
        "KIRKPATRICK, S.; GELATT, C. D.; VECCHI, M. P. Optimization by Simulated Annealing. Science, 1983.",
        "AARTS, E. H. L.; VAN LAARHOVEN, P. J. M. Simulated Annealing: An Introduction. 1989.",
        "TRIKI, E.; COLLETTE, Y.; SIARRY, P. A Theoretical Study on the Behavior of Simulated Annealing. 2005.",
        "ROADS, Curtis. Microsound. MIT Press, 2004."
    ]
    for i, r in enumerate(refs):
        bx = s_b3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.3), Inches(3.2 + i * 1.55), Inches(17.41), Inches(1.2))
        bx.fill.solid()
        bx.fill.fore_color.rgb = SURFACE_COLOR
        bx.line.color.rgb = BORDER_COLOR
        bx.line.width = Pt(1.5)
        p = bx.text_frame.paragraphs[0]
        p.text = r
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT_COLOR
    add_footer(s_b3, "03", is_appendix=True)

    prs.save(output_path)
    print(f"Deck enxuto gerado com sucesso em: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "docs/apresentacao-mochila-tempera-simulada.pptx"
    create_tight_mochila_deck(out)
